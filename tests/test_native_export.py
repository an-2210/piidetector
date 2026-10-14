import io

import docx
import fitz
import pytest
from pptx import Presentation

from src.p5.native_export import redact_native_file

PAIRS = [("Rahul Sharma", "[PERSON_1]"), ("rahul@example.com", "[EMAIL_1]")]


def _docx_bytes():
    document = docx.Document()
    document.add_paragraph("Contact Rahul Sharma at rahul@example.com today.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Owner: Rahul Sharma"
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _pptx_bytes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Rahul Sharma"
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = "mail rahul@example.com"
    slide.notes_slide.notes_text_frame.text = "Rahul Sharma notes"
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def _pdf_bytes():
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Contact Rahul Sharma at rahul@example.com")
    out = document.tobytes()
    document.close()
    return out


def test_docx_stays_docx_and_removes_values():
    result = redact_native_file(_docx_bytes(), ".docx", PAIRS)
    assert result["suffix"] == ".docx"
    reopened = docx.Document(io.BytesIO(result["bytes"]))
    text = "\n".join(p.text for p in reopened.paragraphs) + reopened.tables[0].cell(0, 0).text
    assert "Rahul Sharma" not in text and "rahul@example.com" not in text
    assert "[PERSON_1]" in text and "[EMAIL_1]" in text
    assert result["unlocated_tokens"] == []


def test_pptx_stays_pptx_including_notes():
    result = redact_native_file(_pptx_bytes(), ".pptx", PAIRS)
    reopened = Presentation(io.BytesIO(result["bytes"]))
    slide = reopened.slides[0]
    texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    texts.append(slide.notes_slide.notes_text_frame.text)
    joined = "\n".join(texts)
    assert "Rahul Sharma" not in joined and "rahul@example.com" not in joined
    assert "[PERSON_1]" in joined
    assert result["unlocated_tokens"] == []


def test_pdf_stays_pdf_and_removes_values():
    result = redact_native_file(_pdf_bytes(), ".pdf", PAIRS)
    reopened = fitz.open(stream=result["bytes"], filetype="pdf")
    text = reopened[0].get_text()
    reopened.close()
    assert "Rahul Sharma" not in text and "rahul@example.com" not in text
    assert "[PERSON_1]" in text
    assert result["unlocated_tokens"] == []


def test_value_not_in_file_is_reported_as_unlocated():
    result = redact_native_file(_docx_bytes(), ".docx", [("Jane Doe", "[PERSON_2]")])
    assert result["unlocated_tokens"] == ["[PERSON_2]"]


def test_unsupported_format_is_rejected():
    with pytest.raises(ValueError):
        redact_native_file(b"", ".txt", PAIRS)


def test_preview_text_shows_redacted_docx_and_pptx():
    from src.p5.native_export import preview_text

    docx_out = redact_native_file(_docx_bytes(), ".docx", PAIRS)["bytes"]
    pptx_out = redact_native_file(_pptx_bytes(), ".pptx", PAIRS)["bytes"]
    assert "[PERSON_1]" in preview_text(docx_out, ".docx")
    assert "Rahul Sharma" not in preview_text(pptx_out, ".pptx")


def test_pdf_preview_renders_png_pages():
    from src.p5.native_export import render_pdf_preview

    pages = render_pdf_preview(redact_native_file(_pdf_bytes(), ".pdf", PAIRS)["bytes"], max_pages=2)
    assert len(pages) == 1
    assert pages[0].startswith(b"\x89PNG")

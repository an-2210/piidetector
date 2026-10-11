
"""Pytest fixtures for generating synthetic test documents."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Tuple

import docx
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont
import pptx
from pptx.util import Inches as PptxInches
import pymupdf
import pytest

from src.evaluation.metrics import GroundTruthStructure


@pytest.fixture
def sample_pdf_with_ground_truth(
    tmp_path: Path,
) -> Tuple[Path, GroundTruthStructure]:
    """Create a synthetic PDF with headings, paragraphs, and a vector table."""
    file_path = tmp_path / "threat_report.pdf"

    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)

    page.insert_text(
        (50, 50), "Threat Intelligence Briefing", fontsize=16
    )
    page.insert_text(
        (50, 80),
        "An advanced persistent threat campaign was detected targeting critical infrastructure.",
        fontsize=10,
    )
    page.insert_text(
        (50, 105),
        "Threat actors leveraged spear-phishing attachments to gain initial network access.",
        fontsize=10,
    )

    table_data = [
        ["Adversary", "Target Sector", "Severity"],
        ["APT-29", "Energy Grid", "Critical"],
        ["FIN7", "Financial Services", "High"],
        ["Lazarus", "Cryptocurrency", "High"],
    ]

    rects = [
        (50, 150, 180, 180),
        (180, 150, 350, 180),
        (350, 150, 500, 180),
        (50, 180, 180, 210),
        (180, 180, 350, 210),
        (350, 180, 500, 210),
        (50, 210, 180, 240),
        (180, 210, 350, 240),
        (350, 210, 500, 240),
        (50, 240, 180, 270),
        (180, 240, 350, 270),
        (350, 240, 500, 270),
    ]

    index = 0
    for row in table_data:
        for cell in row:
            rect = rects[index]
            page.draw_rect(pymupdf.Rect(rect), color=(0, 0, 0), width=1)
            page.insert_text(
                (rect[0] + 5, rect[1] + 18), cell, fontsize=9
            )
            index += 1

    page.insert_text(
        (50, 300),
        "Recommended mitigation includes immediate credential revocation and MFA enforcement.",
        fontsize=10,
    )

    doc.save(str(file_path))
    doc.close()

    ground_truth = GroundTruthStructure(
        expected_headings=["Threat Intelligence Briefing"],
        expected_paragraphs=[
            "An advanced persistent threat campaign was detected targeting critical infrastructure.",
            "Threat actors leveraged spear-phishing attachments to gain initial network access.",
            "Recommended mitigation includes immediate credential revocation and MFA enforcement.",
        ],
        expected_tables=[table_data],
    )

    return file_path, ground_truth


@pytest.fixture
def multi_page_ocr_pdf(tmp_path: Path) -> Path:
    """Create a hybrid PDF with native text and an image-only scanned page."""
    file_path = tmp_path / "hybrid_scanned.pdf"

    doc = pymupdf.open()

    # Page 1: Normal native PDF text.
    page_one = doc.new_page(width=595, height=842)
    page_one.insert_text(
        (50, 50),
        "Standard native text on page one.",
        fontsize=12,
    )

    # Page 2: Raster image containing text, with no native text layer.
    page_two = doc.new_page(width=595, height=842)

    image = Image.new("RGB", (1200, 500), color="white")
    draw = ImageDraw.Draw(image)

    try:
        font = ImageFont.truetype("arial.ttf", 48)
    except OSError:
        font = ImageFont.load_default(size=48)

    draw.text(
        (40, 80),
        "Employee ID: EMP-40123",
        fill="black",
        font=font,
    )
    draw.text(
        (40, 180),
        "Account Status: Active",
        fill="black",
        font=font,
    )

    image_buffer = io.BytesIO()
    image.save(image_buffer, format="PNG")

    page_two.insert_image(
        pymupdf.Rect(50, 50, 545, 450),
        stream=image_buffer.getvalue(),
    )

    doc.save(str(file_path))
    doc.close()

    return file_path


@pytest.fixture
def sample_docx_with_ground_truth(
    tmp_path: Path,
) -> Tuple[Path, GroundTruthStructure]:
    """Create a DOCX with headers, footers, headings, paragraphs, and a table."""
    file_path = tmp_path / "incident_response.docx"

    document = docx.Document()

    section = document.sections[0]
    section.header.paragraphs[0].text = (
        "CONFIDENTIAL - SOC INVESTIGATION"
    )
    section.footer.paragraphs[0].text = (
        "Incident Tracking ID: IR-2026-99"
    )

    document.add_heading("Security Incident Analysis", level=1)
    document.add_paragraph(
        "Suspicious outbound C2 communication was observed from host 10.0.4.15."
    )

    table_data = [
        ["Timestamp", "Source IP", "Destination IP", "Action"],
        ["14:02:11", "10.0.4.15", "198.51.100.24", "Blocked"],
        ["14:03:45", "10.0.4.15", "203.0.113.88", "Terminated"],
    ]

    table = document.add_table(
        rows=len(table_data),
        cols=len(table_data[0]),
    )

    for row_index, row in enumerate(table_data):
        for column_index, value in enumerate(row):
            table.cell(row_index, column_index).text = value

    document.add_heading("Containment Protocols", level=2)
    document.add_paragraph(
        "Network isolation applied to infected workload. Forensic snapshot taken."
    )

    document.save(str(file_path))

    ground_truth = GroundTruthStructure(
        expected_headings=[
            "Security Incident Analysis",
            "Containment Protocols",
        ],
        expected_paragraphs=[
            "Suspicious outbound C2 communication was observed from host 10.0.4.15.",
            "Network isolation applied to infected workload. Forensic snapshot taken.",
        ],
        expected_tables=[table_data],
        expected_headers_footers=[
            "CONFIDENTIAL - SOC INVESTIGATION",
            "Incident Tracking ID: IR-2026-99",
        ],
    )

    return file_path, ground_truth


@pytest.fixture
def sample_pptx_with_ground_truth(
    tmp_path: Path,
) -> Tuple[Path, GroundTruthStructure]:
    """Create a presentation with titles, tables, grouped shapes, and notes."""
    file_path = tmp_path / "threat_briefing.pptx"

    presentation = pptx.Presentation()

    # Slide 1: Title and subtitle.
    slide_one = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )
    slide_one.shapes.title.text = (
        "Cyber Threat Intelligence Briefing"
    )
    slide_one.placeholders[1].text = "Quarterly Strategic Review"

    # Slide 2: Title, table, and speaker notes.
    slide_two = presentation.slides.add_slide(
        presentation.slide_layouts[5]
    )
    slide_two.shapes.title.text = "Vulnerability Prioritization"

    table_data = [
        ["CVE ID", "CVSS Score", "Exploited in Wild"],
        ["CVE-2026-1044", "9.8", "Yes"],
        ["CVE-2026-8812", "7.5", "No"],
    ]

    table_shape = slide_two.shapes.add_table(
        len(table_data),
        len(table_data[0]),
        PptxInches(1),
        PptxInches(2),
        PptxInches(6),
        PptxInches(2),
    )

    for row_index, row in enumerate(table_data):
        for column_index, value in enumerate(row):
            table_shape.table.cell(
                row_index, column_index
            ).text = value

    slide_two.notes_slide.notes_text_frame.text = (
        "Emphasize immediate patching of CVE-2026-1044."
    )

    # Slide 3: Grouped shapes.
    slide_three = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )
    group = slide_three.shapes.add_group_shape()

    shape_one = group.shapes.add_shape(
        pptx.enum.shapes.MSO_SHAPE.RECTANGLE,
        PptxInches(1),
        PptxInches(1),
        PptxInches(3),
        PptxInches(1),
    )
    shape_one.text_frame.text = "Grouped Component Alpha"

    shape_two = group.shapes.add_shape(
        pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE,
        PptxInches(1),
        PptxInches(2.5),
        PptxInches(3),
        PptxInches(1),
    )
    shape_two.text_frame.text = "Grouped Component Beta"

    presentation.save(str(file_path))

    ground_truth = GroundTruthStructure(
        expected_headings=[
            "Cyber Threat Intelligence Briefing",
            "Vulnerability Prioritization",
        ],
        expected_paragraphs=[
            "Quarterly Strategic Review",
            "Grouped Component Alpha",
            "Grouped Component Beta",
            "Emphasize immediate patching of CVE-2026-1044.",
        ],
        expected_tables=[table_data],
        expected_slides=[
            "Cyber Threat Intelligence Briefing",
            "Vulnerability Prioritization",
        ],
    )

    return file_path, ground_truth

@pytest.fixture
def docx_with_ocr_image(tmp_path: Path) -> Path:
    """Create a DOCX containing an image with readable employee information."""
    file_path = tmp_path / "employee_scan.docx"

    image = Image.new("RGB", (1400, 400), color="white")
    draw = ImageDraw.Draw(image)

    try:
        font = ImageFont.truetype("arial.ttf", 64)
    except OSError:
        font = ImageFont.load_default(size=64)

    draw.text(
        (40, 80),
        "Employee ID: EMP-40123",
        fill="black",
        font=font,
    )
    draw.text(
        (40, 200),
        "Account Status: Active",
        fill="black",
        font=font,
    )

    image_buffer = io.BytesIO()
    image.save(image_buffer, format="PNG")
    image_buffer.seek(0)

    document = docx.Document()
    document.add_heading("Employee Record", level=1)
    document.add_picture(image_buffer, width=Inches(6))
    document.save(str(file_path))

    return file_path


@pytest.fixture
def pptx_with_ocr_image(tmp_path: Path) -> Path:
    """Create a PPTX containing an image with readable account information."""
    file_path = tmp_path / "account_scan.pptx"

    image = Image.new("RGB", (1400, 400), color="white")
    draw = ImageDraw.Draw(image)

    try:
        font = ImageFont.truetype("arial.ttf", 64)
    except OSError:
        font = ImageFont.load_default(size=64)

    draw.text(
        (40, 80),
        "Account Number: ACC-12345",
        fill="black",
        font=font,
    )
    draw.text(
        (40, 200),
        "Account Status: Active",
        fill="black",
        font=font,
    )

    image_buffer = io.BytesIO()
    image.save(image_buffer, format="PNG")
    image_buffer.seek(0)

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(
        presentation.slide_layouts[6]
    )
    slide.shapes.add_picture(
        image_buffer,
        PptxInches(1),
        PptxInches(1),
        width=PptxInches(8),
    )
    presentation.save(str(file_path))

    return file_path

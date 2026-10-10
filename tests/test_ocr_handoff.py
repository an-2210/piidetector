"""Tests for OCR handoff flagging for downstream P2 pipeline."""

from pathlib import Path
from PIL import Image
import pptx
from pptx.util import Inches
import pymupdf
import pytest
from src.extractors.pdf_extractor import PDFExtractor
from src.extractors.pptx_extractor import PptxExtractor
from src.pipeline import extract_document


def test_ocr_handoff_multipage_pdf(multi_page_ocr_pdf):
    result = extract_document(multi_page_ocr_pdf)

    # Page 1 has text, Page 2 has only raster image
    assert len(result.ocr_needed_pages) == 1
    flag = result.ocr_needed_pages[0]

    assert flag.page_number == 2
    assert flag.reason == "scanned_page_no_text"
    assert flag.has_images is True
    assert flag.char_count == 0

    # Ensure flag serializes to machine-readable dictionary
    flag_dict = flag.to_dict()
    assert flag_dict["page_number"] == 2
    assert flag_dict["reason"] == "scanned_page_no_text"
    assert flag_dict["has_images"] is True


def test_no_ocr_flag_for_rich_text_doc(sample_pdf_with_ground_truth):
    pdf_path, _ = sample_pdf_with_ground_truth
    result = extract_document(pdf_path)
    assert len(result.ocr_needed_pages) == 0, "Rich native text should not trigger OCR flag"


def test_ocr_flag_image_only_pptx(tmp_path):
    """Verify that an image-only slide with no text is flagged for downstream P2."""
    pptx_path = tmp_path / "image_slide.pptx"
    prs = pptx.Presentation()

    # Create slide with blank layout
    s = prs.slides.add_slide(prs.slide_layouts[6])

    # Save a temporary image and add to slide
    img_path = tmp_path / "temp_slide_img.png"
    Image.new("RGB", (200, 200), color="blue").save(str(img_path))
    s.shapes.add_picture(str(img_path), Inches(1), Inches(1), Inches(4), Inches(4))

    prs.save(str(pptx_path))

    extractor = PptxExtractor()
    result = extractor.extract(pptx_path)

    assert len(result.ocr_needed_pages) == 1
    flag = result.ocr_needed_pages[0]
    assert flag.slide_number == 1
    assert flag.reason == "image_only_slide"
    assert flag.has_images is True

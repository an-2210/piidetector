"""Tests for python-pptx based PptxExtractor."""

from pathlib import Path
import pptx
from pptx.util import Inches
import pytest
from src.extractors.pptx_extractor import PptxExtractor
from src.models.text_block import BlockType, SourceType


def test_pptx_extraction_basic(sample_pptx_with_ground_truth):
    pptx_path, gt = sample_pptx_with_ground_truth
    extractor = PptxExtractor()
    result = extractor.extract(pptx_path)

    assert result.source_file == pptx_path.name
    assert result.source_type == SourceType.PPTX.value
    assert result.total_blocks > 0
    assert len(result.errors) == 0


def test_pptx_slide_titles_and_numbers(sample_pptx_with_ground_truth):
    pptx_path, gt = sample_pptx_with_ground_truth
    extractor = PptxExtractor()
    result = extractor.extract(pptx_path)

    title_blocks = [b for b in result.blocks if b.metadata.get("is_slide_title")]
    titles = [b.text for b in title_blocks]

    assert "Cyber Threat Intelligence Briefing" in titles
    assert "Vulnerability Prioritization" in titles

    for b in result.blocks:
        assert b.slide_number is not None
        assert b.slide_number in (1, 2, 3)


def test_pptx_table_extraction(sample_pptx_with_ground_truth):
    pptx_path, gt = sample_pptx_with_ground_truth
    extractor = PptxExtractor()
    result = extractor.extract(pptx_path)

    table_blocks = [b for b in result.blocks if b.block_type == BlockType.TABLE.value]
    assert len(table_blocks) == 1, "Expected 1 table in PPTX"

    tbl = table_blocks[0]
    assert tbl.slide_number == 2
    assert tbl.content is not None
    assert len(tbl.content) == 3  # 3 rows
    assert len(tbl.content[0]) == 3  # 3 cols
    assert "CVE-2026-1044" in tbl.text
    assert "9.8" in tbl.text


def test_pptx_speaker_notes(sample_pptx_with_ground_truth):
    pptx_path, gt = sample_pptx_with_ground_truth
    extractor = PptxExtractor(include_speaker_notes=True)
    result = extractor.extract(pptx_path)

    note_blocks = [b for b in result.blocks if b.metadata.get("is_speaker_notes")]
    assert len(note_blocks) == 1
    assert "CVE-2026-1044" in note_blocks[0].text
    assert note_blocks[0].slide_number == 2


def test_pptx_grouped_shapes(sample_pptx_with_ground_truth):
    pptx_path, gt = sample_pptx_with_ground_truth
    extractor = PptxExtractor()
    result = extractor.extract(pptx_path)

    slide3_blocks = [b for b in result.blocks if b.slide_number == 3]
    slide3_texts = [b.text for b in slide3_blocks]

    assert any("Grouped Component Alpha" in t for t in slide3_texts)
    assert any("Grouped Component Beta" in t for t in slide3_texts)


def test_pptx_shapes_without_text(tmp_path):
    """Ensure shapes without text (e.g. geometric shapes, lines) do not crash the extractor."""
    pptx_file = tmp_path / "shapes_no_text.pptx"
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_shape(
        pptx.enum.shapes.MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(1),
        Inches(1),
        Inches(2),
        Inches(2),
    )
    prs.save(str(pptx_file))

    extractor = PptxExtractor()
    result = extractor.extract(pptx_file)
    assert len(result.errors) == 0
    assert result.total_blocks == 0

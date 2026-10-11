"""Generate the synthetic evaluation documents in eval/samples/.

All names, emails, IDs and addresses are invented. The documents exist only to
score P4. Run with the project venv:  .venv/bin/python eval/samples/make_samples.py
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

import docx
import pymupdf
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

HERE = Path(__file__).resolve().parent
FONT_PATH = "/System/Library/Fonts/Supplemental/Arial.ttf"


def make_prose_docx() -> None:
    document = docx.Document()
    document.add_heading("16.5 Insider Activity", level=1)
    document.add_paragraph(
        "Michael T. Achebe, Senior Analyst at Northwind Logistics Ltd, raised the alert on 3 March. "
        "Reach him at michael.achebe@northwind-logistics.com or EMP-40123. "
        "Ticket RSK-2024-0031 covers the loss, estimated at $12,500. "
        "Policy effective date 01/04/2025 applies."
    )
    document.add_paragraph(
        "The shipment was rerouted to 14 Harbour View Road, Lagos, before reaching the depot on Marina Street, Ibadan."
    )
    document.add_heading("16.8 Third-Party Exposure", level=1)
    document.add_paragraph(
        "The vendor Globex Ltd shared logs with Dr. Priya Raman (EMP-77001). "
        "Acme Corp reviewed them in March 2026. "
        "Escalation was copied to M. Achebe and Michael Achebe."
    )
    document.add_paragraph(
        "Grace Achebe, Head of Audit, signed off the review. Contact priya.raman@globex-ltd.com for follow-up."
    )
    document.save(str(HERE / "prose_case.docx"))


def make_roster_docx() -> None:
    document = docx.Document()
    document.add_heading("Appendix A: Access roster", level=1)
    rows = [
        ["Name", "Employee", "Email", "Office"],
        ["M. T. Achebe", "EMP-40123", "m.achebe@northwind-logistics.com", "14 Harbour View Road, Lagos"],
        ["Grace Achebe", "EMP-40987", "grace.achebe@northwind-logistics.com", "7 Admiralty Way, Lekki"],
        ["Priya Raman", "EMP-77001", "priya.raman@globex-ltd.com", "9 Broad Street, Lagos Island"],
        ["Acme Corp", "vendor", "ops@acme-corp.com", "RSK-2024-0031"],
    ]
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.rows[r].cells[c].text = value
    document.add_paragraph("Totals: USD 4,000 approved for access tooling.")
    document.save(str(HERE / "roster_table.docx"))


def make_briefing_pptx() -> None:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]

    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Insider Risk Briefing"
    slide.placeholders[1].text = "Confidential"

    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Escalations"
    slide.placeholders[1].text = (
        "Michael Achebe escalated case RSK-2024-0031 with the board.\n"
        "Vendor: Globex Ltd, 9 Broad Street, Lagos Island"
    )

    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "Contacts"
    slide.placeholders[1].text = (
        "Contact Priya Raman (EMP-77001), priya.raman@globex-ltd.com\n"
        "Budget $250,000 USD approved"
    )
    presentation.save(str(HERE / "briefing.pptx"))


def make_structured_conflicts_docx() -> None:
    """Structured identifiers (what P3 owns) and deliberate overlaps for fusion."""
    document = docx.Document()
    document.add_heading("Case annex: identifiers and conflicts", level=1)
    lines = [
        "Applicant: Michael T. Achebe reviewed the case with the board.",
        "Applicant: Grace Achebe, Head of Audit, approved the exception.",
        "Home address: 9 Marina Road, Ikoyi.",
        "Email: priya.raman@globex-ltd.com",
        "Follow-up was sent to priya.raman@globex-ltd.com for the Lagos office details.",
        "SSN: 123-45-6789 is on file. Card ending 4421 and account ending 9087 were flagged.",
        "PAN 4111 1111 1111 1111 was declined. PESEL: 80010112345. TIN: 12-3456789. DOB: 14/03/1988.",
        "Directory ID DIR-88421 and merchant IDs MER-IN-449120 and MER-PH-220011 were linked.",
    ]
    for line in lines:
        document.add_paragraph(line)
    document.save(str(HERE / "structured_conflicts.docx"))


def make_decoys_memo_docx() -> None:
    """Negatives that look like PII, plus a few real positives so precision is measured."""
    document = docx.Document()
    document.add_heading("Risk Committee minutes", level=1)
    lines = [
        "Escalation went to the Risk Committee in March, after the review.",
        "Dana Holt, Chief Risk Officer, approved the memo. Registered office: 22 Allen Avenue, Ikeja.",
        "Achebe Holdings Ltd pays the fee for the Lagos office.",
        "Several items were logged. Following the audit, Policy GRP-POL-LAGOS-01 applies.",
        "In March 2026 the RSK-2024-0031 file was closed. The Monday review used USD 18,000.",
        "Finance and Legal reviewed the file. Head of Compliance signed off. The SSN and DOB fields were checked by the Lagos Operations Team.",
    ]
    for line in lines:
        document.add_paragraph(line)
    document.save(str(HERE / "decoys_memo.docx"))


def _render_scan_page() -> bytes:
    """Render text as a JPEG scan, so the PDF has no text layer and P2 must OCR it."""
    lines = [
        "CONFIDENTIAL INTAKE FORM",
        "Applicant: Michael T. Achebe",
        "Email: michael.achebe@northwind-logistics.com",
        "Employee ID: EMP-40123",
        "Home address: 14 Harbour View Road, Lagos",
        "Sponsor: Acme Corp",
        "Reference: RSK-2024-0031",
        "Policy effective date 01/04/2025",
    ]
    image = Image.new("RGB", (1700, 1100), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(FONT_PATH, 40)
    y = 120
    for line in lines:
        draw.text((120, y), line, fill="black", font=font)
        y += 90
    buffer = BytesIO()
    image.save(buffer, format="JPEG", quality=90)
    return buffer.getvalue()


def make_scanned_pdf() -> None:
    document = pymupdf.open()
    page = document.new_page(width=612, height=396)
    page.insert_image(page.rect, stream=_render_scan_page())
    document.save(str(HERE / "scanned_intake.pdf"))
    document.close()


if __name__ == "__main__":
    make_prose_docx()
    make_roster_docx()
    make_briefing_pptx()
    make_scanned_pdf()
    make_structured_conflicts_docx()
    make_decoys_memo_docx()
    for name in ("prose_case.docx", "roster_table.docx", "briefing.pptx", "scanned_intake.pdf", "structured_conflicts.docx", "decoys_memo.docx"):
        print("wrote", name, (HERE / name).stat().st_size, "bytes")

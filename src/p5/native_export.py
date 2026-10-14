"""Write a redacted copy of an uploaded document in its original format.

Input PDF, DOCX and PPTX files produce PDF, DOCX and PPTX output. Each detected
value is replaced by its token. Values that cannot be located as text in the
file (for example text inside scanned images) are listed by token in the report.
"""
from __future__ import annotations

import io
from typing import Any, Dict, Iterable, Iterator, List, Sequence, Tuple

import fitz  # PyMuPDF
from docx import Document as load_docx
from pptx import Presentation

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".pptx"}


def _longest_first(pairs: Iterable[Tuple[str, str]]) -> List[Tuple[str, str]]:
    cleaned = {(value, token) for value, token in pairs if value}
    return sorted(cleaned, key=lambda pair: len(pair[0]), reverse=True)


def _replace_in_runs(runs: Sequence[Any], pairs: List[Tuple[str, str]], applied: set) -> None:
    """Replace inside each run first to keep formatting; merge only for values split across runs."""
    for run in runs:
        for value, token in pairs:
            if value in run.text:
                run.text = run.text.replace(value, token)
                applied.add(token)
    full = "".join(run.text for run in runs)
    leftover = [(value, token) for value, token in pairs if value in full]
    if leftover and runs:
        for value, token in leftover:
            full = full.replace(value, token)
            applied.add(token)
        runs[0].text = full
        for run in runs[1:]:
            run.text = ""


def _table_paragraphs_docx(table) -> Iterator[Any]:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _table_paragraphs_docx(nested)


def _docx_paragraphs(document) -> Iterator[Any]:
    yield from document.paragraphs
    for table in document.tables:
        yield from _table_paragraphs_docx(table)
    for section in document.sections:
        for part in (section.header, section.footer):
            yield from part.paragraphs
            for table in part.tables:
                yield from _table_paragraphs_docx(table)


def _redact_docx(data: bytes, pairs: List[Tuple[str, str]]) -> Tuple[bytes, set, List[str]]:
    document = load_docx(io.BytesIO(data))
    applied: set = set()
    for paragraph in list(_docx_paragraphs(document)):
        _replace_in_runs(paragraph.runs, pairs, applied)
    texts = [paragraph.text for paragraph in _docx_paragraphs(document)]
    out = io.BytesIO()
    document.save(out)
    return out.getvalue(), applied, _leaked_tokens(texts, pairs)


def _pptx_text_frames(shapes) -> Iterator[Any]:
    for shape in shapes:
        if hasattr(shape, "shapes"):  # group shape
            yield from _pptx_text_frames(shape.shapes)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape.text_frame
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def _pptx_all_text_frames(presentation) -> List[Any]:
    frames = []
    for slide in presentation.slides:
        frames.extend(_pptx_text_frames(slide.shapes))
        if slide.has_notes_slide:
            frames.append(slide.notes_slide.notes_text_frame)
    return [frame for frame in frames if frame is not None]


def _redact_pptx(data: bytes, pairs: List[Tuple[str, str]]) -> Tuple[bytes, set, List[str]]:
    presentation = Presentation(io.BytesIO(data))
    applied: set = set()
    for frame in _pptx_all_text_frames(presentation):
        for paragraph in frame.paragraphs:
            _replace_in_runs(paragraph.runs, pairs, applied)
    texts = [
        paragraph.text
        for frame in _pptx_all_text_frames(presentation)
        for paragraph in frame.paragraphs
    ]
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue(), applied, _leaked_tokens(texts, pairs)


def _redact_pdf(data: bytes, pairs: List[Tuple[str, str]]) -> Tuple[bytes, set, List[str]]:
    document = fitz.open(stream=data, filetype="pdf")
    applied: set = set()
    for page in document:
        for value, token in pairs:
            rects = page.search_for(value)
            for rect in rects:
                page.add_redact_annot(rect, text=token, fill=(0, 0, 0), text_color=(1, 1, 1))
            if rects:
                applied.add(token)
        page.apply_redactions()
    texts = [page.get_text() for page in document]
    out = document.tobytes(garbage=3, deflate=True)
    document.close()
    return out, applied, _leaked_tokens(texts, pairs)


def _leaked_tokens(texts: Sequence[str], pairs: List[Tuple[str, str]]) -> List[str]:
    """Tokens whose original value still appears in the output text. Reports tokens only, never values."""
    joined = "\n".join(texts)
    return sorted({token for value, token in pairs if value in joined})


def redact_native_file(data: bytes, suffix: str, replacements: Iterable[Tuple[str, str]]) -> Dict[str, Any]:
    """Return {"bytes", "suffix", "unlocated_tokens"} for a redacted copy of ``data``.

    ``replacements`` holds (original value, token) pairs from the redaction step.
    Raises ValueError for unsupported formats.
    """
    suffix = suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Redacted export is not available for {suffix or 'this'} files")
    pairs = _longest_first(replacements)
    if suffix == ".pdf":
        out, applied, leaks = _redact_pdf(data, pairs)
    elif suffix == ".docx":
        out, applied, leaks = _redact_docx(data, pairs)
    else:
        out, applied, leaks = _redact_pptx(data, pairs)
    # A value that was never found in the file is unlocated, even if no text is left behind.
    unlocated = {token for _, token in pairs if token not in applied} | set(leaks)
    return {"bytes": out, "suffix": suffix, "unlocated_tokens": sorted(unlocated)}


def render_pdf_preview(data: bytes, max_pages: int = 2, dpi: int = 110) -> List[bytes]:
    """PNG bytes for the first pages of a redacted PDF, for an on-screen check before download."""
    document = fitz.open(stream=data, filetype="pdf")
    try:
        pages = []
        for index in range(min(max_pages, document.page_count)):
            pixmap = document[index].get_pixmap(dpi=dpi)
            pages.append(pixmap.tobytes("png"))
        return pages
    finally:
        document.close()


def preview_text(data: bytes, suffix: str, max_chars: int = 4000) -> str:
    """Plain text of a redacted DOCX or PPTX, truncated for display."""
    suffix = suffix.lower()
    if suffix == ".docx":
        document = load_docx(io.BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in _docx_paragraphs(document))
    elif suffix == ".pptx":
        presentation = Presentation(io.BytesIO(data))
        parts = []
        for number, slide in enumerate(presentation.slides, start=1):
            frames = [frame.text for frame in _slide_frames(slide)]
            parts.append(f"Slide {number}\n" + "\n".join(frames))
        text = "\n\n".join(parts)
    else:
        raise ValueError(f"No text preview for {suffix} files")
    return text[:max_chars] + ("\n…" if len(text) > max_chars else "")


def _slide_frames(slide) -> List[Any]:
    frames = list(_pptx_text_frames(slide.shapes))
    if slide.has_notes_slide:
        frames.append(slide.notes_slide.notes_text_frame)
    return frames

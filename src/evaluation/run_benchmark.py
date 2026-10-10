"""Standalone structure retention benchmark script.

Runs extraction across synthetic PDF, DOCX, and PPTX documents with known
ground truths, computing retention scores and printing a comprehensive report.
"""

from __future__ import annotations

import io
import json
import os
from pathlib import Path
import sys
import tempfile

# Ensure project root is on sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

import docx
from PIL import Image
import pptx
from pptx.util import Inches as PptxInches
import pymupdf

from src.evaluation.metrics import GroundTruthStructure, StructureRetentionEvaluator
from src.pipeline import extract_document


def create_benchmark_pdf(dir_path: Path) -> tuple[Path, GroundTruthStructure]:
    pdf_path = dir_path / "benchmark_threat_report.pdf"
    doc = pymupdf.open()
    page = doc.new_page(width=595, height=842)

    page.insert_text((50, 50), "Executive Cybersecurity Threat Summary", fontsize=15)
    page.insert_text(
        (50, 80),
        "Threat actors initiated high-volume distributed denial-of-service attacks against API gateways.",
        fontsize=10,
    )
    page.insert_text(
        (50, 105),
        "Automated edge mitigation filtering rules successfully suppressed 94% of malicious traffic volume.",
        fontsize=10,
    )

    table_data = [
        ["Attack Vector", "Peak Volume", "Status"],
        ["SYN Flood", "420 Gbps", "Mitigated"],
        ["DNS Amplification", "180 Gbps", "Blocked"],
        ["HTTP/2 Rapid Reset", "12 Mpps", "Rate-Limited"],
    ]

    rects = [
        (50, 140, 190, 165), (190, 140, 340, 165), (340, 140, 500, 165),
        (50, 165, 190, 190), (190, 165, 340, 190), (340, 165, 500, 190),
        (50, 190, 190, 215), (190, 190, 340, 215), (340, 190, 500, 215),
        (50, 215, 190, 240), (190, 215, 340, 240), (340, 215, 500, 240),
    ]

    idx = 0
    for row in table_data:
        for cell in row:
            r = rects[idx]
            page.draw_rect(pymupdf.Rect(r), color=(0, 0, 0), width=1)
            page.insert_text((r[0] + 5, r[1] + 16), cell, fontsize=9)
            idx += 1

    page.insert_text(
        (50, 270),
        "Continuous monitoring and rate-limiting safeguards remain enforced.",
        fontsize=10,
    )

    doc.save(str(pdf_path))
    doc.close()

    gt = GroundTruthStructure(
        expected_headings=["Executive Cybersecurity Threat Summary"],
        expected_paragraphs=[
            "Threat actors initiated high-volume distributed denial-of-service attacks against API gateways.",
            "Automated edge mitigation filtering rules successfully suppressed 94% of malicious traffic volume.",
            "Continuous monitoring and rate-limiting safeguards remain enforced.",
        ],
        expected_tables=[table_data],
    )
    return pdf_path, gt


def create_benchmark_docx(dir_path: Path) -> tuple[Path, GroundTruthStructure]:
    docx_path = dir_path / "benchmark_incident_response.docx"
    doc = docx.Document()

    section = doc.sections[0]
    section.header.paragraphs[0].text = "OPTIV INCIDENT INVESTIGATION"
    section.footer.paragraphs[0].text = "Document Classification: Internal"

    doc.add_heading("Threat Assessment Report", level=1)
    doc.add_paragraph("Anomalous lateral movement was observed across segment 192.168.10.0/24.")

    table_data = [
        ["Asset", "Risk Score", "Vulnerability"],
        ["Domain Controller 01", "9.2", "Zerologon"],
        ["File Server 03", "7.8", "SMBv1 Active"],
    ]
    tbl = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            tbl.cell(r_idx, c_idx).text = val

    doc.add_heading("Remediation Roadmap", level=2)
    doc.add_paragraph("All vulnerable legacy services disabled and Kerberos auditing enabled.")

    doc.save(str(docx_path))

    gt = GroundTruthStructure(
        expected_headings=["Threat Assessment Report", "Remediation Roadmap"],
        expected_paragraphs=[
            "Anomalous lateral movement was observed across segment 192.168.10.0/24.",
            "All vulnerable legacy services disabled and Kerberos auditing enabled.",
        ],
        expected_tables=[table_data],
        expected_headers_footers=[
            "OPTIV INCIDENT INVESTIGATION",
            "Document Classification: Internal",
        ],
    )
    return docx_path, gt


def create_benchmark_pptx(dir_path: Path) -> tuple[Path, GroundTruthStructure]:
    pptx_path = dir_path / "benchmark_briefing.pptx"
    prs = pptx.Presentation()

    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Enterprise Threat Briefing"
    s1.placeholders[1].text = "Adversary Tactics and Trends"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "MITRE ATT&CK Matrix Mapping"

    table_data = [
        ["Tactic", "Technique ID", "Detection Status"],
        ["Execution", "T1059.001", "Covered"],
        ["Persistence", "T1547.001", "Covered"],
    ]
    tbl_shape = s2.shapes.add_table(
        len(table_data),
        len(table_data[0]),
        PptxInches(1),
        PptxInches(2),
        PptxInches(6),
        PptxInches(2),
    )
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            tbl_shape.table.cell(r_idx, c_idx).text = val

    s2.notes_slide.notes_text_frame.text = "Speaker note: verify Sigma rules for T1059.001."

    prs.save(str(pptx_path))

    gt = GroundTruthStructure(
        expected_headings=["Enterprise Threat Briefing", "MITRE ATT&CK Matrix Mapping"],
        expected_paragraphs=[
            "Adversary Tactics and Trends",
            "Speaker note: verify Sigma rules for T1059.001.",
        ],
        expected_tables=[table_data],
        expected_slides=["Enterprise Threat Briefing", "MITRE ATT&CK Matrix Mapping"],
    )
    return pptx_path, gt


def run_benchmark():
    evaluator = StructureRetentionEvaluator()
    print("=" * 80)
    print("  P1 NATIVE FORMAT EXTRACTION: STRUCTURE RETENTION BENCHMARK SUITE")
    print("=" * 80)

    with tempfile.TemporaryDirectory() as td:
        tmp_dir = Path(td)

        # 1. PDF
        pdf_path, pdf_gt = create_benchmark_pdf(tmp_dir)
        pdf_res = extract_document(pdf_path)
        pdf_rep = evaluator.evaluate(pdf_res, pdf_gt)

        # 2. DOCX
        docx_path, docx_gt = create_benchmark_docx(tmp_dir)
        docx_res = extract_document(docx_path)
        docx_rep = evaluator.evaluate(docx_res, docx_gt)

        # 3. PPTX
        pptx_path, pptx_gt = create_benchmark_pptx(tmp_dir)
        pptx_res = extract_document(pptx_path)
        pptx_rep = evaluator.evaluate(pptx_res, pptx_gt)

        results = [
            ("PDF", pdf_res, pdf_rep),
            ("DOCX", docx_res, docx_rep),
            ("PPTX", pptx_res, pptx_rep),
        ]

        print(f"{'Format':<8} | {'Blocks':<8} | {'Paragraphs':<12} | {'Tables':<10} | {'Headings':<10} | {'Composite':<12} | {'Status':<10}")
        print("-" * 80)
        all_passed = True
        for fmt, res, rep in results:
            passed = rep.composite_score >= 0.80
            if not passed:
                all_passed = False
            status = "PASS (>=80%)" if passed else "FAIL"
            print(
                f"{fmt:<8} | "
                f"{res.total_blocks:<8} | "
                f"{rep.paragraph_retention * 100:>8.1f}%   | "
                f"{rep.table_retention * 100:>7.1f}%  | "
                f"{rep.heading_retention * 100:>7.1f}%  | "
                f"{rep.composite_score * 100:>9.2f}%   | "
                f"{status:<10}"
            )

        print("-" * 80)
        avg_score = sum(r[2].composite_score for r in results) / len(results)
        print(f"Overall Average Structure Retention: {avg_score * 100:.2f}% (Target: >= 80.00%)")
        print(f"Benchmark Result: {'ALL FORMATS PASSED TARGET' if all_passed else 'FAILED'}")
        print("=" * 80)


if __name__ == "__main__":
    run_benchmark()

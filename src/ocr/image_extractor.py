
"""Extract image content from PDF, DOCX and PPTX documents."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from typing import Any, Optional

import pymupdf
from PIL import Image
from docx import Document
from pptx import Presentation


@dataclass
class ImageItem:
    """An image with its location in the original document."""

    source_file: str
    source_type: str
    image: Image.Image
    page_number: Optional[int] = None
    slide_number: Optional[int] = None
    metadata: dict[str, Any] = field(default_factory=dict)


def _open_image(data: bytes) -> Optional[Image.Image]:
    """Decode image bytes safely."""
    try:
        with Image.open(BytesIO(data)) as image:
            image.load()
            return image.convert("RGB")
    except Exception:
        return None


def extract_image_items(
    file_path: str | Path,
    extraction_result=None,
    render_dpi: int = 200,
) -> list[ImageItem]:
    """Extract embedded images and render PDF pages flagged for OCR.

    P1's ExtractionResult is optional. If provided, its OCR flags identify
    pages to render. Embedded images are also extracted independently.
    """
    path = Path(file_path)
    extension = path.suffix.lower()
    items: list[ImageItem] = []

    if extension == ".pdf":
        flagged_pages = {
            flag.page_number
            for flag in (
                extraction_result.ocr_needed_pages
                if extraction_result is not None
                else []
            )
            if flag.page_number is not None
        }

        with pymupdf.open(str(path)) as document:
            for page_index, page in enumerate(document):
                page_number = page_index + 1

                # Extract embedded images, including those on mixed-content pages.
                for image_info in page.get_images(full=True):
                    xref = image_info[0]
                    image_data = document.extract_image(xref)
                    decoded = _open_image(image_data["image"])

                    if decoded is not None:
                        items.append(
                            ImageItem(
                                source_file=path.name,
                                source_type="PDF",
                                image=decoded,
                                page_number=page_number,
                                metadata={
                                    "kind": "embedded_image",
                                    "xref": xref,
                                    "width": decoded.width,
                                    "height": decoded.height,
                                    "rectangles": [
                                        tuple(float(v) for v in rect)
                                        for rect in page.get_image_rects(xref)
                                    ],
                                },
                            )
                        )

                # Render flagged pages to recover text from full-page scans.
                if page_number in flagged_pages:
                    scale = render_dpi / 72
                    pixmap = page.get_pixmap(
                        matrix=pymupdf.Matrix(scale, scale),
                        alpha=False,
                    )
                    rendered = _open_image(pixmap.tobytes("png"))

                    if rendered is not None:
                        items.append(
                            ImageItem(
                                source_file=path.name,
                                source_type="PDF",
                                image=rendered,
                                page_number=page_number,
                                metadata={
                                    "kind": "rendered_page",
                                    "dpi": render_dpi,
                                    "width": rendered.width,
                                    "height": rendered.height,
                                },
                            )
                        )

    elif extension == ".docx":
        document = Document(str(path))

        for relationship_id, relationship in document.part.rels.items():
            if "image" not in relationship.reltype:
                continue

            image = _open_image(relationship.target_part.blob)
            if image is None:
                continue

            items.append(
                ImageItem(
                    source_file=path.name,
                    source_type="DOCX",
                    image=image,
                    metadata={
                        "kind": "embedded_image",
                        "relationship_id": relationship_id,
                        "width": image.width,
                        "height": image.height,
                    },
                )
            )

    elif extension == ".pptx":
        presentation = Presentation(str(path))

        for slide_number, slide in enumerate(
            presentation.slides, start=1
        ):
            for shape in slide.shapes:
                if shape.shape_type != 13:  # PICTURE
                    continue

                image = _open_image(shape.image.blob)
                if image is None:
                    continue

                items.append(
                    ImageItem(
                        source_file=path.name,
                        source_type="PPTX",
                        image=image,
                        slide_number=slide_number,
                        metadata={
                            "kind": "embedded_image",
                            "shape_name": shape.name,
                            "width": image.width,
                            "height": image.height,
                            "bbox": [
                                shape.left,
                                shape.top,
                                shape.left + shape.width,
                                shape.top + shape.height,
                            ],
                        },
                    )
                )

    else:
        raise ValueError(f"Unsupported file format: {extension}")

    return items

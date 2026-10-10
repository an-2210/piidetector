"""Native PPTX text and table extraction using python-pptx.

Extracts slide titles, body shapes, nested grouped shapes, tables, and speaker
notes with spatial ordering and OCR flagging for image-only slides.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.extractors.base import BaseExtractor
from src.models.text_block import (
    BlockType,
    ExtractionResult,
    OCRPageFlag,
    SourceType,
    TextBlock,
)
from src.utils.helpers import clean_text, format_table_as_markdown

logger = logging.getLogger(__name__)


class PptxExtractor(BaseExtractor):
    """Extracts native text, tables, grouped shapes, and notes from PPTX presentations."""

    supported_extensions = (".pptx",)

    def __init__(self, include_speaker_notes: bool = True) -> None:
        """Initialize PptxExtractor.

        Args:
            include_speaker_notes: Whether to extract speaker notes from slides.
        """
        self.include_speaker_notes = include_speaker_notes

    def _extract_shape_contents(
        self,
        shape: Any,
        source_file: str,
        slide_num: int,
        slide_title: Optional[str],
    ) -> List[Tuple[Tuple[int, int], TextBlock]]:
        """Recursively extract TextBlocks from a single shape or grouped shapes.

        Returns a list of tuples ((top, left), TextBlock) for spatial sorting.
        """
        extracted: List[Tuple[Tuple[int, int], TextBlock]] = []

        top = getattr(shape, "top", 0) or 0
        left = getattr(shape, "left", 0) or 0
        spatial_key = (int(top), int(left))

        # 1. Grouped shapes - recursive traversal
        if hasattr(shape, "shapes"):
            for child_shape in shape.shapes:
                extracted.extend(
                    self._extract_shape_contents(
                        child_shape, source_file, slide_num, slide_title
                    )
                )
            return extracted

        # 2. Table shape
        if getattr(shape, "has_table", False):
            table = shape.table
            rows_data: List[List[str]] = []
            for row in table.rows:
                row_cells: List[str] = []
                for cell in row.cells:
                    cell_text = "\n".join(
                        p.text.strip() for p in cell.text_frame.paragraphs if p.text.strip()
                    )
                    row_cells.append(cell_text)
                rows_data.append(row_cells)

            if rows_data and any(any(c for c in row) for row in rows_data):
                md_table = format_table_as_markdown(rows_data)
                tbl_block = TextBlock(
                    source_file=source_file,
                    source_type=SourceType.PPTX.value,
                    block_type=BlockType.TABLE.value,
                    text=md_table,
                    content=rows_data,
                    slide_number=slide_num,
                    section_or_heading=slide_title,
                    metadata={
                        "shape_id": getattr(shape, "shape_id", None),
                        "row_count": len(rows_data),
                        "col_count": len(rows_data[0]) if rows_data else 0,
                        "top": top,
                        "left": left,
                    },
                )
                extracted.append((spatial_key, tbl_block))
            return extracted

        # 3. Text Frame shape (Text box, placeholder, callout, etc.)
        if getattr(shape, "has_text_frame", False):
            text_frame = shape.text_frame
            text_runs: List[str] = []
            for p in text_frame.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    text_runs.append(p_text)

            full_text = "\n".join(text_runs)
            if full_text:
                shape_block = TextBlock(
                    source_file=source_file,
                    source_type=SourceType.PPTX.value,
                    block_type=BlockType.SLIDE.value,
                    text=full_text,
                    slide_number=slide_num,
                    section_or_heading=slide_title,
                    metadata={
                        "shape_id": getattr(shape, "shape_id", None),
                        "shape_name": getattr(shape, "name", ""),
                        "top": top,
                        "left": left,
                    },
                )
                extracted.append((spatial_key, shape_block))

        return extracted

    def extract(
        self,
        file_path: Union[str, Path],
        include_speaker_notes: Optional[bool] = None,
        **kwargs: Any,
    ) -> ExtractionResult:
        """Extract structured TextBlocks from a PPTX presentation.

        Args:
            file_path: Path to the .pptx presentation.
            include_speaker_notes: Override flag for speaker notes extraction.

        Returns:
            ExtractionResult containing structured slide blocks, tables, and metadata.
        """
        path = self.validate_file(file_path)
        do_notes = include_speaker_notes if include_speaker_notes is not None else self.include_speaker_notes

        result = ExtractionResult(
            source_file=path.name,
            source_type=SourceType.PPTX.value,
        )

        try:
            prs = pptx.Presentation(str(path))

            result.metadata = {
                "slide_count": len(prs.slides),
                "slide_width": getattr(prs, "slide_width", 0),
                "slide_height": getattr(prs, "slide_height", 0),
            }

            reading_order_idx = 0

            for slide_idx, slide in enumerate(prs.slides):
                slide_num = slide_idx + 1

                # 1. Identify slide title if present
                slide_title: Optional[str] = None
                title_shape_id: Optional[int] = None

                if slide.shapes.title and slide.shapes.title.has_text_frame:
                    raw_title = slide.shapes.title.text_frame.text.strip()
                    if raw_title:
                        slide_title = raw_title
                        title_shape_id = getattr(slide.shapes.title, "shape_id", None)

                # Track slide content elements
                slide_elements: List[Tuple[Tuple[int, int], TextBlock]] = []
                image_shapes_count = 0

                # 2. Extract title as first slide block if available
                if slide_title:
                    title_block = TextBlock(
                        source_file=path.name,
                        source_type=SourceType.PPTX.value,
                        block_type=BlockType.SLIDE.value,
                        text=slide_title,
                        slide_number=slide_num,
                        section_or_heading=slide_title,
                        reading_order=reading_order_idx,
                        metadata={
                            "is_slide_title": True,
                            "shape_id": title_shape_id,
                        },
                    )
                    result.blocks.append(title_block)
                    reading_order_idx += 1

                # 3. Extract all other shapes on slide
                for shape in slide.shapes:
                    if title_shape_id is not None and getattr(shape, "shape_id", None) == title_shape_id:
                        continue  # Skip title shape already processed

                    # Track image shapes for potential OCR handoff flag
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        image_shapes_count += 1

                    extracted_blocks = self._extract_shape_contents(
                        shape, path.name, slide_num, slide_title
                    )
                    slide_elements.extend(extracted_blocks)

                # Sort slide elements by reading order: (top, left)
                slide_elements.sort(key=lambda item: item[0])

                for _, block in slide_elements:
                    block.reading_order = reading_order_idx
                    reading_order_idx += 1
                    result.blocks.append(block)

                # 4. Check for image-only or low-text slide needing OCR handoff
                slide_blocks = [
                    b for b in result.blocks
                    if b.slide_number == slide_num and not b.metadata.get("is_speaker_notes")
                ]
                total_slide_text = "".join(b.text for b in slide_blocks)
                if len(total_slide_text.strip()) == 0 and image_shapes_count > 0:
                    result.ocr_needed_pages.append(
                        OCRPageFlag(
                            source_file=path.name,
                            slide_number=slide_num,
                            reason="image_only_slide",
                            char_count=0,
                            has_images=True,
                            metadata={"image_shapes_count": image_shapes_count},
                        )
                    )

                # 5. Extract Speaker Notes if requested
                if do_notes and slide.has_notes_slide:
                    try:
                        notes_frame = slide.notes_slide.notes_text_frame
                        if notes_frame:
                            notes_text = clean_text(notes_frame.text)
                            if notes_text:
                                notes_block = TextBlock(
                                    source_file=path.name,
                                    source_type=SourceType.PPTX.value,
                                    block_type=BlockType.SLIDE.value,
                                    text=notes_text,
                                    slide_number=slide_num,
                                    section_or_heading=slide_title,
                                    reading_order=reading_order_idx,
                                    metadata={
                                        "is_speaker_notes": True,
                                    },
                                )
                                result.blocks.append(notes_block)
                                reading_order_idx += 1
                    except Exception as note_err:
                        logger.debug("Non-fatal speaker note extraction issue: %s", note_err)

            result.total_blocks = len(result.blocks)

        except Exception as e:
            logger.error("PPTX extraction error for %s: %s", path.name, e)
            result.errors.append(f"PPTX extraction error: {e}")

        return result
